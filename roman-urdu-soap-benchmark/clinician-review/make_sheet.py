"""Build a blinded review sheet for a physiotherapist.

50 notes: all 40 Roman Urdu notes and 10 English notes (5 per model, chosen at
random with a fixed seed). Notes are shuffled and given codes R01..R50. The
sheet shows the English transcript as the reference (the facts are the same in
both languages), so the reviewer does not see the model, the language of the
input or the automatic check results.

Writes review_sheet.xlsx (for the reviewer) and review_key.csv (code -> note).
The reviewer must not open review_key.csv.
"""
from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
BENCH = HERE.parent
sys.path.insert(0, str(BENCH))
from run import load_cases  # noqa: E402

SEED = 42
EN_PER_MODEL = 5
LABELS = ["wrong_region", "side_problem", "missing_critical", "other_error"]
HEADERS = {
    "wrong_region": "Wrong or missing body region (Y/N)",
    "side_problem": "Side wrong or missing (Y/N)",
    "missing_critical": "Red flag or other critical fact missing (Y/N)",
    "other_error": "Other error, e.g. invented or wrong fact (Y/N)",
}
INSTRUCTIONS = [
    "Thank you for reviewing these notes.",
    "Each row is one SOAP note drafted by software from the consultation on the left.",
    "Read the consultation, then the note, and mark each column Y (yes) or N (no):",
    "1. Wrong or missing body region: the note names the wrong body part, adds one that was not discussed, or does not name the main one.",
    "2. Side wrong or missing: the consultation names left or right and the note gives the wrong side or no side. Mark N if no side was mentioned.",
    "3. Red flag or other critical fact missing: a red-flag answer, medicine, referral or urgent action from the consultation is not in the note.",
    "4. Other error: anything else a physiotherapist would correct, such as a symptom or detail that was never said, or a wrong number.",
    "Use the Comment column for anything you want to explain. Please judge each note on its own.",
    "Please do not open review_key.csv until you have finished.",
]


def select_notes(notes: List[Dict[str, object]], seed: int = SEED) -> List[Dict[str, object]]:
    rng = random.Random(seed)
    chosen = [n for n in notes if n["lang"] == "ru"]
    for model in sorted({str(n["model"]) for n in notes}):
        en = [n for n in notes if n["lang"] == "en" and n["model"] == model]
        chosen += rng.sample(en, EN_PER_MODEL)
    rng.shuffle(chosen)
    return chosen


def build(notes_path: Path, out_dir: Path) -> List[Dict[str, str]]:
    notes = [json.loads(l) for l in notes_path.read_text(encoding="utf-8").splitlines() if l.strip()]
    cases = {c["id"]: c for c in load_cases()}
    chosen = select_notes(notes)

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, 1):
        ins.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True)
    ins.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Review")
    cols = ["Code", "Consultation (reference)", "Drafted note"] + [HEADERS[k] for k in LABELS] + ["Comment"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(yn)
    fill = PatternFill("solid", fgColor="FFF2CC")

    key: List[Dict[str, str]] = []
    for i, n in enumerate(chosen, 1):
        code = f"R{i:02d}"
        ref, note = str(cases[n["case"]]["en"]), str(n["note"])
        ws.append([code, ref, note] + [""] * (len(LABELS) + 1))
        r = ws.max_row
        for col in range(1, len(cols) + 1):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(4, 4 + len(LABELS)):
            yn.add(ws.cell(row=r, column=col))
            ws.cell(row=r, column=col).fill = fill
        lines = max(ref.count("\n"), note.count("\n")) + max(len(ref), len(note)) // 60
        ws.row_dimensions[r].height = min(409, 15 * (lines + 2))
        key.append({"code": code, "note_id": str(n["note_id"]), "model": str(n["model"]), "lang": str(n["lang"])})
    for letter, width in zip("ABCDEFGH", [7, 60, 70, 16, 16, 18, 18, 30]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "B2"

    out_dir.mkdir(exist_ok=True)
    wb.save(out_dir / "review_sheet.xlsx")
    with (out_dir / "review_key.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(key[0]))
        w.writeheader()
        w.writerows(key)
    return key


if __name__ == "__main__":
    k = build(BENCH / "results" / "notes.jsonl", HERE)
    print(f"{len(k)} notes: {sum(r['lang'] == 'ru' for r in k)} Roman Urdu, {sum(r['lang'] == 'en' for r in k)} English")
