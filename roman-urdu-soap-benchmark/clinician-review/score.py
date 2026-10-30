"""Compare the physiotherapist's labels with the automatic checks.

Input: the filled review sheet (.xlsx, or .csv exported from it), review_key.csv
and ../results/per_note.csv.
Output: labels.csv (her labels joined to the automatic checks) and
agreement.csv (per check: counts, raw agreement, Cohen's kappa, and precision
and recall of the automatic check with her labels as the reference).

Usage: python score.py filled_review_sheet.xlsx
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path
from typing import Dict, List, Optional

HERE = Path(__file__).resolve().parent
PER_NOTE = HERE.parent / "results" / "per_note.csv"
LABELS = ["wrong_region", "side_problem", "missing_critical", "other_error"]


def read_sheet(path: Path) -> Dict[str, Dict[str, str]]:
    """Return {code: {label: 'Y'|'N', 'comment': str}} from the Review sheet."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig") as fh:
            rows = [list(r) for r in csv.reader(fh)]
    else:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True)["Review"]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    out: Dict[str, Dict[str, str]] = {}
    for r in rows[1:]:
        if not r or not str(r[0]).strip():
            continue
        vals = [str(v).strip().upper()[:1] for v in r[3:3 + len(LABELS)]]
        if any(v not in ("Y", "N") for v in vals):
            raise ValueError(f"{r[0]}: every check must be Y or N, got {vals}")
        out[str(r[0]).strip()] = dict(zip(LABELS, vals), comment=str(r[3 + len(LABELS)]) if len(r) > 3 + len(LABELS) else "")
    return out


def auto_labels(row: Dict[str, str]) -> Dict[str, bool]:
    """The automatic checks, mapped to the reviewer's columns."""
    return {
        "wrong_region": row["main_region_present"] == "False" or bool(row["unsupported_regions"]),
        "side_problem": row["side_present"] == "False" or row["wrong_side"] == "True",
        "missing_critical": int(row["critical_missing"]) > 0,
        "gate": row["gate_flag"] == "True",
    }


def kappa(a: List[bool], b: List[bool]) -> Optional[float]:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return None if pe == 1 else (po - pe) / (1 - pe)


def compare(auto: List[bool], human: List[bool], name: str, subset: str) -> Dict[str, object]:
    tp = sum(a and h for a, h in zip(auto, human))
    fp = sum(a and not h for a, h in zip(auto, human))
    fn = sum(h and not a for a, h in zip(auto, human))
    tn = len(auto) - tp - fp - fn
    k = kappa(auto, human)
    return {
        "check": name, "subset": subset, "n": len(auto),
        "both_yes": tp, "auto_only": fp, "reviewer_only": fn, "both_no": tn,
        "agreement": round((tp + tn) / len(auto), 3),
        "kappa": "" if k is None else round(k, 3),
        "precision": "" if tp + fp == 0 else round(tp / (tp + fp), 3),
        "recall": "" if tp + fn == 0 else round(tp / (tp + fn), 3),
    }


def score(sheet: Path, key_path: Path = HERE / "review_key.csv", per_note: Path = PER_NOTE,
          out: Path = HERE) -> List[Dict[str, object]]:
    human = read_sheet(sheet)
    with key_path.open(encoding="utf-8") as fh:
        key = list(csv.DictReader(fh))
    with per_note.open(encoding="utf-8") as fh:
        checks = {f"{r['case']}-{r['lang']}|{r['model']}": r for r in csv.DictReader(fh)}
    missing = [k["code"] for k in key if k["code"] not in human]
    if missing:
        raise ValueError(f"not yet reviewed: {', '.join(missing)}")

    joined: List[Dict[str, object]] = []
    for k in key:
        a = auto_labels(checks[f"{k['note_id']}|{k['model']}"])
        h = human[k["code"]]
        row: Dict[str, object] = dict(k)
        for lab in LABELS:
            row[f"reviewer_{lab}"] = h[lab] == "Y"
        row["reviewer_any_error"] = any(h[lab] == "Y" for lab in LABELS)
        for lab, v in a.items():
            row[f"auto_{lab}"] = v
        row["reviewer_comment"] = h["comment"]
        joined.append(row)

    results: List[Dict[str, object]] = []
    pairs = [("wrong_region", "wrong_region"), ("side_problem", "side_problem"),
             ("missing_critical", "missing_critical"), ("gate", "any_error")]
    for subset, rows in (("all", joined), ("roman_urdu", [r for r in joined if r["lang"] == "ru"])):
        for auto_name, hum_name in pairs:
            results.append(compare([bool(r[f"auto_{auto_name}"]) for r in rows],
                                   [bool(r[f"reviewer_{hum_name}"]) for r in rows], auto_name, subset))
        results.append({"check": "reviewer_other_error", "subset": subset, "n": len(rows),
                        "both_yes": sum(bool(r["reviewer_other_error"]) for r in rows)})

    for name, rows_out in (("labels.csv", joined), ("agreement.csv", results)):
        fields: List[str] = []
        for r in rows_out:
            fields += [f for f in r if f not in fields]
        with (out / name).open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=fields)
            w.writeheader()
            w.writerows(rows_out)
    return results


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: python score.py filled_review_sheet.xlsx")
    for r in score(Path(sys.argv[1])):
        print(r)
